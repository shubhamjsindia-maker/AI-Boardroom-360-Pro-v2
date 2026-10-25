"""Preview-grade downloadable historical analysis exports."""
from io import BytesIO
from .financial import STATEMENTS,ratios

def excel_report(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    wb=Workbook();cover=wb.active;cover.title='About';cover.append(['AI Boardroom 360 Pro — Historical Financial Review']);cover.append(['Company',data['company']]);cover.append(['Reporting Basis',data.get('basis','')]);cover.append(['Unit','INR lakh']);cover.append(['Source',data.get('source','')]);cover.append(['Status','DRAFT / UNVERIFIED — not audited'])
    years=list(data['periods'])
    for name,items in STATEMENTS.items():
        sh=wb.create_sheet({'Profit & Loss':'Profit and Loss','Balance Sheet':'Balance Sheet','Cash Flow':'Cash Flow','Changes in Equity':'Changes in Equity'}[name]);sh.append(['Financial Line Item']+years)
        for key,label in items: sh.append([label]+[data['periods'][y].get(key) for y in years])
        sh.freeze_panes='B2';sh.column_dimensions['A'].width=48
        for c in sh[1]: c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='183858')
        for idx in range(2,len(years)+2): sh.column_dimensions[__import__('openpyxl').utils.get_column_letter(idx)].width=19
    ratiosheet=wb.create_sheet('Financial Ratios');ratiosheet.append(['Ratio']+years)
    keys=list(ratios(data['periods'][years[-1]]))
    for key in keys: ratiosheet.append([key]+[float(ratios(data['periods'][y]).get(key)) if ratios(data['periods'][y]).get(key) is not None else None for y in years])
    ratiosheet.column_dimensions['A'].width=42
    stream=BytesIO();wb.save(stream);return stream.getvalue()

def pdf_report(data,findings):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer,PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    output=BytesIO();doc=SimpleDocTemplate(output,pagesize=(595,842),leftMargin=35,rightMargin=35,topMargin=40,bottomMargin=40)
    styles=getSampleStyleSheet();elements=[Paragraph('AI Boardroom 360 Pro | Historical Financial Review',styles['Title']),Paragraph(data['company'],styles['Heading2']),Paragraph('DRAFT / UNVERIFIED — NOT AUDITED | INR lakh',styles['Normal']),Spacer(1,15)]
    years=list(data['periods'])
    for name,items in STATEMENTS.items():
        elements.append(Paragraph(name,styles['Heading2']))
        rows=[['Financial statement line']+years]
        for k,label in items:
            rows.append([label]+[str(data['periods'][y].get(k,'—')) for y in years])
        widths=[240]+[min(85,280/len(years))]*len(years)
        table=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#183858')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTSIZE',(0,0),(-1,-1),7.2),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,0),(-1,0),0.6,colors.grey)]))
        elements.extend([table,Spacer(1,14)])
    elements.extend([Paragraph('Validation Summary',styles['Heading2']),Paragraph(f'{len(findings)} critical findings. Source figures require independent verification.',styles['Normal'])])
    doc.build(elements);return output.getvalue()
